import json
import re
from pathlib import Path

import openpyxl
import pytest

import import_viewpoints as iv

VP_HEADER = [
    "ID",
    "種別",
    "観点",
    "対応テストケース",
    "前提",
    "操作",
    "期待結果",
    "出典",
    "統合元・備考",
]
CODE_HEADER = [
    "独自コード",
    "エラーID(msgid)",
    "ステータス",
    "対応観点ID",
    "対応ケースID",
    "概要",
    "検証手段",
    "照合",
    "備考",
]

SPEC_MD = """# spec

## 2. テストケース

#### S1-01 【前提作成】トークン
#### S5-01 入力不備
#### S5-02 ファイル
#### S9-09 どの観点にもない

## 3. 観点 ID → ケース ID 対応表

| 観点 ID | ケース ID | ケース数 |
| ------- | --------- | -------- |
| BD-01 | S5-01 | 1 |
| AU-10 | S5-01、S5-02 | 2 |
| UI-05 | S5-02 | 1 |

## 4. 要確認
"""


def make_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "テスト概要"
    ws = wb.create_sheet("テスト観点")

    def vp(*cells):
        ws.append([None, *cells])

    ws.append([])
    vp("各テスト観点の内容")
    vp(*VP_HEADER)
    vp("分類・独自コード・応答形式")
    vp("BD-01", "BD", "独自コードが `error` の先頭に付与される", "S5-01", "Direct のクライアント。")
    ws.cell(ws.max_row, 7, "EP2 で送る")
    ws.cell(ws.max_row, 8, "応答が `X: msg` 形式")
    ws.cell(ws.max_row, 9, "基本 4.2")
    ws.append([])
    vp("認証・認可")
    vp("AU-10", "AU", "スコープ不足は 403", "S5-01, S5-02", "前提", "操作。", "期待。")
    ws.cell(ws.max_row, 9, "基本 3.4, 4.3（U-13）、詳細 2.4")
    vp("ログ方針")
    vp("UI-05", "UI", "内部エラーは 501", "S5-02", "前提", "操作", "期待", "独自の文書")

    ws2 = wb.create_sheet("独自コード対応表")

    def code(*cells):
        ws2.append([None, *cells])

    ws2.append([])
    code("説明")
    code(*CODE_HEADER)
    code(3201, "INTERNAL_SERVER_ERROR", 501, "UI-05", "S5-02", "`msg` 501", "モック", "明示")
    code(1201, "ACTIVITY_SCOPE_INSUFFICIENT", 403, "AU-10, BD-01", "S5-02, S5-01", "`m` 403")
    ws2.cell(ws2.max_row, 8, "実環境（設定変更が必要）")
    ws2.cell(ws2.max_row, 9, "帯のみ")
    wb.create_sheet("旧テスト観点")
    wb.save(path)


@pytest.fixture
def inputs(tmp_path):
    xlsx = tmp_path / "in.xlsx"
    make_workbook(xlsx)
    md = tmp_path / "spec.md"
    md.write_text(SPEC_MD, encoding="utf-8")
    return xlsx, md


def test_parse_sources_splits_documents_sections_and_notes():
    refs = iv.parse_sources("基本 3.4, 4.3（DB 上は推定。U-13）、詳細 2.4（注意事項、処理仕様 6）")
    assert refs == [
        {"kind": "design", "ref": "basic_design.md v1.4 §3.4, §4.3", "note": "DB 上は推定。U-13"},
        {"kind": "design", "ref": "detailed_design.md v1.3 §2.4", "note": "注意事項、処理仕様 6"},
    ]


def test_parse_sources_falls_back_to_other():
    assert iv.parse_sources("独自の文書") == [{"kind": "other", "ref": "独自の文書"}]
    assert iv.parse_sources("基本 abc") == [{"kind": "other", "ref": "基本 abc"}]
    assert iv.parse_sources("")[0]["kind"] == "other"


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("plain text 日本語", "plain text 日本語"),
        ("has: colon", '"has: colon"'),
        ("`leading backtick`", '"`leading backtick`"'),
        ("mid `tick` ok", "mid `tick` ok"),
        ("1411", '"1411"'),
        ("true", '"true"'),
        ("", '""'),
        ('say "hi" # not a comment', '"say \\"hi\\" # not a comment"'),
        ("ends with colon:", '"ends with colon:"'),
        ("a#b", "a#b"),
        ("- dash", '"- dash"'),
    ],
)
def test_yaml_scalar_matches_hub_writer_rules(value, expected):
    assert iv.yaml_scalar(value) == expected


def test_yaml_scalar_rejects_control_characters():
    with pytest.raises(ValueError):
        iv.yaml_scalar("a\nb")


def test_risk_rule(inputs):
    book = iv.read_workbook(inputs[0])
    by_id = {v.id: v for v in book.viewpoints}
    status = {"UI-05": 501}
    assert iv.risk_of(by_id["AU-10"], {}) == "high"  # AU family
    assert iv.risk_of(by_id["UI-05"], status) == "high"  # maps to a 5xx code
    assert iv.risk_of(by_id["UI-05"], {}) == "low"  # section 'ログ方針'
    assert iv.risk_of(by_id["BD-01"], {}) == "medium"


def test_read_workbook_skips_section_rows_and_tracks_sections(inputs):
    book = iv.read_workbook(inputs[0])
    assert [v.id for v in book.viewpoints] == ["BD-01", "AU-10", "UI-05"]
    assert [v.section for v in book.viewpoints] == [
        "分類・独自コード・応答形式",
        "認証・認可",
        "ログ方針",
    ]
    assert book.viewpoints[1].cases == ["S5-01", "S5-02"]
    assert [c.code for c in book.codes] == ["3201", "1201"]
    assert book.codes[0].status == 501


def test_viewpoint_yaml_is_canonical(inputs):
    book = iv.read_workbook(inputs[0])
    entity = iv.viewpoint_entity(book.viewpoints[1], {})
    assert iv.viewpoint_yaml(entity) == (
        "id: VP-SW-AU-10\n"
        "title: スコープ不足は 403\n"
        "rationale: 前提：前提。操作：操作。期待結果：期待。\n"
        "source:\n"
        "  - kind: design\n"
        "    ref: basic_design.md v1.4 §3.4, §4.3\n"
        "    note: U-13\n"
        "  - kind: design\n"
        "    ref: detailed_design.md v1.3 §2.4\n"
        "risk: high\n"
        "parents: []\n"
    )


def test_end_to_end_outputs_are_reproducible(inputs, tmp_path):
    xlsx, md = inputs
    first, second = tmp_path / "a", tmp_path / "b"
    iv.run(xlsx, md, first)
    iv.run(xlsx, md, second)
    files_a = sorted(p.relative_to(first) for p in first.rglob("*") if p.is_file())
    files_b = sorted(p.relative_to(second) for p in second.rglob("*") if p.is_file())
    assert files_a == files_b
    assert [str(p) for p in files_a] == [
        "docs/spec-trace.md",
        "fixtures/error-code-map.json",
        "specs/viewpoints/VP-SW-AU-10.yaml",
        "specs/viewpoints/VP-SW-BD-01.yaml",
        "specs/viewpoints/VP-SW-UI-05.yaml",
    ]
    for rel in files_a:
        assert (first / rel).read_bytes() == (second / rel).read_bytes()
    for path in (first / "specs/viewpoints").iterdir():
        assert re.match(r"^id: VP-[A-Z0-9]+(-[A-Z0-9]+)*$", path.read_text().splitlines()[0])


def test_code_map_is_sorted_and_complete(inputs, tmp_path):
    iv.run(*inputs, tmp_path / "out")
    entries = json.loads((tmp_path / "out/fixtures/error-code-map.json").read_text("utf-8"))
    assert [e["code"] for e in entries] == ["1201", "3201"]
    first = entries[0]
    assert list(first) == [
        "code", "msgid", "status", "viewpoints", "cases", "summary",
        "verification", "match", "reachability_note",
    ]  # fmt: skip
    assert first["viewpoints"] == ["VP-SW-AU-10", "VP-SW-BD-01"]
    assert first["cases"] == ["S5-01", "S5-02"]
    assert first["reachability_note"] == "実環境（設定変更が必要）"
    assert entries[1]["reachability_note"] == ""


def test_trace_reports_gaps_and_mismatches(inputs, tmp_path):
    iv.run(*inputs, tmp_path / "out")
    text = (tmp_path / "out/docs/spec-trace.md").read_text("utf-8")
    assert "- S1-01：前提作成ケース（観点に紐づかない）" in text
    assert "- S9-09：**要確認**（観点表から参照されていない）" in text
    # AU-10 lists S5-01、S5-02 on both sides; BD-01 and UI-05 agree as well.
    assert "### 観点表と仕様書 3 章の対応表の食い違い\n\nなし。" in text
    assert "- UI-10：観点表 v4.2 で除外（欠番）。観点表に存在しない（想定どおり）" in text


def test_trace_reports_table_mismatch(inputs, tmp_path):
    xlsx, md = inputs
    md.write_text(SPEC_MD.replace("| UI-05 | S5-02 | 1 |", "| UI-05 | S5-01 | 1 |"), "utf-8")
    iv.run(xlsx, md, tmp_path / "out")
    text = (tmp_path / "out/docs/spec-trace.md").read_text("utf-8")
    assert "- UI-05: 観点表 S5-02 / 仕様書 3 章 S5-01" in text


def test_bad_input_is_reported(tmp_path):
    bad = tmp_path / "bad.xlsx"
    openpyxl.Workbook().save(bad)
    md = tmp_path / "s.md"
    md.write_text(SPEC_MD, encoding="utf-8")
    assert iv.main(["--xlsx", str(bad), "--spec-md", str(md), "--out-root", str(tmp_path)]) == 2
    assert iv.main([]) == 2


# ------------------------------------------------------------------ owner overrides


def _write_overrides(path, **change):
    item = {
        "id": "AU-10",
        "field": "expected",
        "from": "期待",
        "to": "新しい期待（ERROR 行の有無では判定しない）",
        "reason": "オーナー決定",
    }
    item.update(change)
    path.write_text(json.dumps([item], ensure_ascii=False), encoding="utf-8")
    return path


def test_override_replaces_the_cell_and_keeps_the_original_in_the_yaml(inputs, tmp_path):
    xlsx, md = inputs
    overrides = _write_overrides(tmp_path / "o.json")
    out = tmp_path / "out"
    counts = iv.run(xlsx, md, out, overrides)
    assert counts["overrides"] == 1
    text = (out / "specs/viewpoints/VP-SW-AU-10.yaml").read_text(encoding="utf-8")
    assert "期待結果：新しい期待" in text
    assert "kind: other" in text and "viewpoint_overrides.json" in text
    assert "オーナー決定" in text and "原文: 期待" in text
    # untouched viewpoints are identical to a run without overrides
    plain = tmp_path / "plain"
    iv.run(xlsx, md, plain)
    other = "specs/viewpoints/VP-SW-BD-01.yaml"
    assert (out / other).read_bytes() == (plain / other).read_bytes()


def test_override_output_is_reproducible(inputs, tmp_path):
    xlsx, md = inputs
    overrides = _write_overrides(tmp_path / "o.json")
    iv.run(xlsx, md, tmp_path / "a", overrides)
    iv.run(xlsx, md, tmp_path / "b", overrides)
    name = "specs/viewpoints/VP-SW-AU-10.yaml"
    assert (tmp_path / "a" / name).read_bytes() == (tmp_path / "b" / name).read_bytes()


def test_a_stale_override_is_an_error(inputs, tmp_path):
    xlsx, md = inputs
    overrides = _write_overrides(tmp_path / "o.json", **{"from": "どこにもない文"})
    with pytest.raises(iv.ImportError_, match="original text is not in the workbook"):
        iv.run(xlsx, md, tmp_path / "out", overrides)


def test_an_override_already_applied_in_the_workbook_is_obsolete_not_an_error(
    inputs, tmp_path, capsys
):
    xlsx, md = inputs
    overrides = _write_overrides(tmp_path / "o.json", **{"from": "古い文", "to": "期待"})
    counts = iv.run(xlsx, md, tmp_path / "out", overrides)
    assert counts["overrides"] == 0
    assert "obsolete" in capsys.readouterr().err


def test_override_of_an_unknown_viewpoint_or_field_is_an_error(inputs, tmp_path):
    xlsx, md = inputs
    with pytest.raises(iv.ImportError_, match="unknown viewpoint"):
        iv.run(xlsx, md, tmp_path / "o1", _write_overrides(tmp_path / "a.json", id="ZZ-99"))
    with pytest.raises(iv.ImportError_, match="field must be one of"):
        iv.run(xlsx, md, tmp_path / "o2", _write_overrides(tmp_path / "b.json", field="title"))


def test_the_committed_overrides_file_is_well_formed():
    overrides = iv.read_overrides(iv.DEFAULT_OVERRIDES)
    assert {o.id for o in overrides} >= {"BD-60", "EG-19"}
    for override in overrides:
        # the new text states the rule; it must not say that an ERROR line means failure
        assert "合否" in override.replacement
        assert "不合格" not in override.replacement
